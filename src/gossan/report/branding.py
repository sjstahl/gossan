"""The company's look for Word reports: logo, colours, fonts, header, footer, cover page, and a Word template.

Two ways to brand, which can be combined:

* **Settings** (``branding:`` in ``project.yaml``): company, logo, accent colour, fonts, footer text,
  contact, cover page. Gossan applies them to the document it builds.
* **A template**: ``gossan template`` writes a starter ``.docx`` from those settings; open it in Word, change
  the styles, page setup, header and footer to the house style, and point ``branding.template`` at it. The
  report is then built *inside* that file, so its styles, page setup and header and footer are used as they
  are (the template's own text is dropped).

Branding cannot remove the disclaimers: the footer line that marks a report as internal or draft is always
added to every page's footer (the cover page too), after whatever the template's footer holds, and the
banner paragraph is part of the report body.

Not checked here: how Word lays a template out. The tests read the file back (styles, header, footer, fields,
images, alt text) but no Word or LibreOffice was available to open it.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage

from gossan.core.project import Project

HEADER_LOGO_HEIGHT_IN = 0.45
COVER_LOGO_HEIGHT_IN = 1.1
TEMPLATE_NAME = "template.docx"
#: Styles the report writer uses; made when a template does not define them.
NEEDED = {"Title": (22, True), "Heading 1": (15, True), "Heading 2": (12.5, True), "Heading 3": (11, True), "Caption": (8.5, False), "List Bullet": (10, False)}


class BrandingError(Exception):
    """The branding settings cannot be used; the message says what to change."""


@dataclass
class Branding:
    company: str | None = None
    logo: Path | None = None
    accent: str = "2A78D6"  # no leading #
    font: str = "Calibri"
    heading_font: str = "Calibri"
    footer_text: str | None = None
    contact: str | None = None
    template: Path | None = None
    cover_page: bool = False
    client: str | None = None
    project_name: str = ""


def resolve(project: Project) -> Branding:
    """The project's branding with paths made absolute and checked. Raises ``BrandingError`` for a logo or
    template that is missing or unreadable, rather than quietly producing an unbranded client document."""
    cfg = project.config.branding

    def path(value: str | None, what: str) -> Path | None:
        if not value:
            return None
        p = Path(value)
        p = p if p.is_absolute() else project.root / p
        if not p.is_file():
            raise BrandingError(f"branding.{what} {value!r} not found (looked for {p})")
        return p

    logo, template = path(cfg.logo, "logo"), path(cfg.template, "template")
    if logo is not None:
        try:
            with PILImage.open(logo) as im:
                im.verify()
        except Exception as e:
            raise BrandingError(f"branding.logo {cfg.logo!r} is not a readable image (use PNG or JPEG): {e}") from e
    return Branding(
        company=cfg.company, logo=logo, accent=cfg.accent.lstrip("#").upper(), font=cfg.font, heading_font=cfg.heading_font or cfg.font,
        footer_text=cfg.footer_text, contact=cfg.contact, template=template, cover_page=cfg.cover_page, client=project.config.client, project_name=project.config.name,
    )  # fmt: skip


def tint(hex_colour: str, amount: float = 0.85) -> str:
    """The colour mixed with white (``amount`` of the way): the light fill for table headers."""
    r, g, b = (int(hex_colour[i : i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in (r, g, b))


# --- the document ------------------------------------------------------------------------------------------------------------


def _clear_body(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_font(style, name: str) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # a theme font would override the name
            rfonts.attrib.pop(qn(f"w:{attr}"), None)
        rfonts.set(qn("w:eastAsia"), name)


def ensure_styles(doc) -> None:
    """Define the styles the report uses that the document lacks (a template made elsewhere may not have them)."""
    present = {s.name for s in doc.styles}
    for name, (size, bold) in NEEDED.items():
        if name in present:
            continue
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.size = Pt(size)
        style.font.bold = bold
        if name == "Caption":
            style.font.italic = True
        if name.startswith("Heading") or name == "Title":
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.keep_with_next = True


def new_document(b: Branding):
    """A blank document in the company's look: the template emptied of its text, or the default one styled
    from the settings."""
    if b.template is not None:
        try:
            doc = Document(str(b.template))
        except Exception as e:
            raise BrandingError(f"cannot open branding.template {b.template.name}: {e}. Save it from Word as .docx (not .dotx).") from e
        _clear_body(doc)
        ensure_styles(doc)
        return doc
    doc = Document()
    ensure_styles(doc)
    normal = doc.styles["Normal"]
    _set_font(normal, b.font)
    normal.font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        _set_font(style, b.heading_font)
        style.font.color.rgb = RGBColor.from_string(b.accent)
    _set_font(doc.styles["Caption"], b.font)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("555555")
    return doc


def bullet(doc, text_runner) -> None:
    """A bullet paragraph filled by ``text_runner(paragraph)``; without a List Bullet style, a bullet sign instead."""
    p = doc.add_paragraph(style="List Bullet")
    if not _has_numbering(doc.styles["List Bullet"]):  # a style we made, or a template's without a list
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run("•  ")
    text_runner(p)


def _has_numbering(style) -> bool:
    ppr = style.element.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def grid_table(table, colour: str = "A6A6A6") -> None:
    """Thin borders on every cell edge, set directly so it does not depend on a style the template may lack."""
    pr = table._tbl.tblPr
    old = pr.find(qn("w:tblBorders"))
    if old is not None:
        pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    pr.append(borders)


# --- header, footer, cover -----------------------------------------------------------------------------------------------


def _field(paragraph, instruction: str, **font) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        elif kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
    _style_run(run, **font)


def _style_run(run, size: float | None = None, bold: bool = False, colour: str | None = None) -> None:
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour)


def _rule(paragraph, colour: str, where: str = "bottom") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{where}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "8")
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), colour)
    border.append(edge)
    ppr.append(border)


def _picture(run, path: Path, height_in: float, alt: str) -> None:
    run.add_picture(str(path), height=Inches(height_in))
    for doc_pr in run._r.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt)  # alt text for screen readers


def _text_width(section):
    return section.page_width - section.left_margin - section.right_margin


def _has_content(part) -> bool:
    return any(p.text.strip() for p in part.paragraphs) or bool(part._element.xpath(".//w:drawing")) or bool(part.tables)


def _write_header(part, section, b: Branding) -> None:
    if _has_content(part):
        return  # the template's own header stands
    part.is_linked_to_previous = False
    p = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(_text_width(section), WD_TAB_ALIGNMENT.RIGHT)
    if b.company:
        _style_run(p.add_run(b.company), size=11, bold=True, colour=b.accent)
    if b.logo:
        p.add_run("\t")
        _picture(p.add_run(), b.logo, HEADER_LOGO_HEIGHT_IN, f"{b.company or 'Company'} logo")
    if b.company or b.logo:
        _rule(p, b.accent)


def _write_footer(part, b: Branding, line: str, name: str, colour: str = "555555") -> None:
    part.is_linked_to_previous = False
    p = part.add_paragraph() if _has_content(part) else (part.paragraphs[0] if part.paragraphs else part.add_paragraph())
    parts = [line, b.company or "", name]
    _style_run(p.add_run("   |   ".join(x for x in parts if x) + "   |   Page "), size=7, bold=True, colour=colour)
    _field(p, "PAGE", size=7, colour=colour)


def add_furniture(doc, b: Branding, footer_line: str, project_name: str, marked: bool = False) -> None:
    """Header (company and logo, unless the template has one) and footer (``footer_line`` always added) on every section."""
    for section in doc.sections:
        section.different_first_page_header_footer = b.cover_page
        _write_header(section.header, section, b)
        colour = "7B1D12" if marked else "555555"
        _write_footer(section.footer, b, footer_line, project_name, colour)
        if b.cover_page:  # the cover has no header, but it carries the footer and so the marking
            _write_footer(section.first_page_footer, b, footer_line, project_name, colour)


def add_cover(doc, b: Branding, title: str, subtitle: str, banner=None) -> None:
    """A cover page: logo, company, title, subtitle, client, date, contact, then the banner (``banner(doc)``)."""
    if b.logo:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(36)
        _picture(p.add_run(), b.logo, COVER_LOGO_HEIGHT_IN, f"{b.company or 'Company'} logo")
    if b.company:
        _style_run(doc.add_paragraph().add_run(b.company), size=14, bold=True, colour=b.accent)
    doc.add_paragraph(style="Title").add_run(title)
    if subtitle:
        _style_run(doc.add_paragraph().add_run(subtitle), size=10, colour="555555")
    lines = [f"Prepared for {b.client}" if b.client else "", f"{datetime.now(UTC):%d %B %Y}", b.contact or ""]
    for line in (x for x in lines if x):
        _style_run(doc.add_paragraph().add_run(line), size=11)
    if banner:
        banner(doc)
    doc.add_page_break()


def set_properties(doc, b: Branding, title: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.author = b.company or "Gossan"


# --- the starter template ----------------------------------------------------------------------------------------------


def write_starter_template(path: Path, b: Branding) -> Path:
    """A ``.docx`` showing the styles, with the company's header and footer, to be edited in Word and named in ``branding.template``."""
    plain = Branding(**{**b.__dict__, "template": None})
    doc = new_document(plain)
    for section in doc.sections:  # reports add their own footer line (marking and page number) after this
        _write_header(section.header, section, plain)
        _style_run(section.footer.paragraphs[0].add_run(" | ".join(x for x in (b.company, b.contact) if x)), size=8, colour="555555")
    doc.add_paragraph(style="Title").add_run("House style for Gossan reports")
    doc.add_paragraph().add_run(
        "Edit the styles below (Home > Styles, right-click > Modify), the page setup, and the header and footer, then save as .docx. "
        "Point branding.template in project.yaml at this file. Gossan builds each Word report inside it and drops this text; "
        "it adds a footer line with the draft or internal marking and the page number to every page, so leave room for it."
    )
    doc.add_paragraph(style="Heading 1").add_run("Heading 1")
    doc.add_paragraph("Body text (Normal) looks like this. The report's sentences, table cells and notes all use it.")
    doc.add_paragraph(style="Heading 2").add_run("Heading 2")
    bullet(doc, lambda p: p.add_run("A bullet (List Bullet)"))
    doc.add_paragraph(style="Caption").add_run("Figure 1. A caption (Caption), under a map or chart.")
    set_properties(doc, plain, "Gossan report template")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


__all__ = ["Branding", "BrandingError", "add_cover", "add_furniture", "bullet", "grid_table", "new_document", "resolve", "tint", "write_starter_template"]
