"""Minimal report: AOI + MINFILE map and a data-sources section from provenance."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import geopandas as gpd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pyogrio
from docx.shared import Inches

from gossan.core.project import Project
from gossan.core.provenance import ProvenanceRecord
from gossan.report.branding import (
    add_furniture,
    bullet,
    new_document,
    resolve,
    set_properties,
)

MINFILE_LAYER = "minfile"
MAP_NAME = "map.png"
REPORT_NAME = "report.docx"


def read_minfile(project: Project) -> gpd.GeoDataFrame | None:
    """The ``minfile`` layer, or None if harvest has not produced one."""
    if not project.data_gpkg.exists():
        return None
    if MINFILE_LAYER not in pyogrio.list_layers(project.data_gpkg)[:, 0]:
        return None
    return gpd.read_file(project.data_gpkg, layer=MINFILE_LAYER, engine="pyogrio")


def make_map(project: Project) -> Path:
    """Draw the AOI and MINFILE points to ``outputs/map.png``."""
    aoi = project.read_aoi()
    points = read_minfile(project)
    project.outputs_dir.mkdir(exist_ok=True)
    path = project.outputs_dir / MAP_NAME

    fig, ax = plt.subplots(figsize=(7, 7))
    aoi.plot(ax=ax, facecolor="#f4d03f33", edgecolor="#b9770e", linewidth=1.5)
    handles = [plt.Line2D([], [], color="#b9770e", linewidth=1.5, label="AOI")]
    n_points = 0
    if points is not None and not points.empty:
        n_points = len(points)
        points.plot(ax=ax, color="#c0392b", edgecolor="black", markersize=40, zorder=3)
        handles.append(
            plt.Line2D(
                [], [], marker="o", linestyle="", markerfacecolor="#c0392b",
                markeredgecolor="black", label=f"MINFILE occurrence (n={n_points})",
            )
        )
    minx, miny, maxx, maxy = aoi.total_bounds
    pad = max(maxx - minx, maxy - miny) * 0.08
    ax.set_xlim(minx - pad, maxx + pad)
    ax.set_ylim(miny - pad, maxy + pad)
    ax.set_aspect("equal")
    ax.set_title(f"{project.config.name}: AOI and MINFILE occurrences")
    ax.set_xlabel(f"Easting (m), {project.config.crs}")
    ax.set_ylabel("Northing (m)")
    ax.ticklabel_format(style="plain", useOffset=False)
    ax.tick_params(axis="x", labelrotation=30)
    ax.grid(True, linewidth=0.3)
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, -0.17), ncol=2)
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight")
    plt.close(fig)

    project.provenance.log(
        dataset="map",
        source="gossan",
        identifier=f"outputs/{MAP_NAME}",
        url=None,
        license="N/A (derived product)",
        operation="render_map",
        parameters={"crs": project.config.crs, "n_minfile": n_points},
    )
    return path


def data_sources(records: list[ProvenanceRecord]) -> list[ProvenanceRecord]:
    """Latest retrieval per (dataset, identifier), from fetched or manually ingested data."""
    latest: dict[tuple[str, str], ProvenanceRecord] = {}
    for r in records:
        if r.operation not in {"fetch", "manual_ingest"}:
            continue
        key = (r.dataset, r.identifier)
        if key not in latest or r.retrieved_at >= latest[key].retrieved_at:
            latest[key] = r
    return sorted(latest.values(), key=lambda r: (r.dataset, r.identifier))


def build_report(project: Project) -> Path:
    """Write ``outputs/report.docx`` (and the map it embeds)."""
    sources = data_sources(project.provenance.read())  # read before our own records are added
    map_path = make_map(project)
    points = read_minfile(project)
    n_points = 0 if points is None else len(points)

    brand = resolve(project)
    doc = new_document(brand)
    add_furniture(doc, brand, brand.footer_text or 'DRAFT | Machine-generated: for geologist review', project.config.name)
    set_properties(doc, brand, project.config.name)
    doc.add_heading(project.config.name, level=0)
    if project.config.client:
        doc.add_paragraph(f"Client: {project.config.client}")
    doc.add_paragraph(f"Project CRS: {project.config.crs}")
    doc.add_paragraph(f"Generated {datetime.now(UTC):%Y-%m-%d} by Gossan. DRAFT for geologist review.")

    doc.add_heading("Mineral occurrences", level=1)
    if points is None:
        doc.add_paragraph("MINFILE data have not been harvested for this project.")
    else:
        doc.add_paragraph(f"{n_points} MINFILE occurrence(s) fall within the area of interest.")
    doc.add_picture(str(map_path), width=Inches(6))
    doc.add_paragraph("Figure 1. Area of interest and MINFILE mineral occurrences.")

    doc.add_heading("Data Sources", level=1)
    if not sources:
        doc.add_paragraph("No external datasets have been retrieved for this project.")
    for r in sources:
        def fill(p, r=r) -> None:
            p.add_run(f"{r.dataset}").bold = True
            p.add_run(f": {r.source}, {r.identifier}. Retrieved {r.retrieved_at:%Y-%m-%d}. Licence: {r.license}.")
            if r.url:
                p.add_run(f" {r.url}")

        bullet(doc, fill)
    if any("British Columbia" in r.license for r in sources):
        doc.add_paragraph(
            "Contains information licensed under the Open Government Licence - British Columbia."
        )

    project.outputs_dir.mkdir(exist_ok=True)
    path = project.outputs_dir / REPORT_NAME
    doc.save(path)
    project.provenance.log(
        dataset="report",
        source="gossan",
        identifier=f"outputs/{REPORT_NAME}",
        url=None,
        license="N/A (derived product)",
        operation="build_report",
        parameters={"n_sources": len(sources), "n_minfile": n_points},
    )
    return path
