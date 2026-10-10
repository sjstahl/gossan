"""Sample pack: a one-property summary of a lead, the reports available, and before/after
samples of what digitizing them produced.

It lists what exists (from the tenure and the ARIS index) and whether each report PDF is
already in ``raw/aris_pdfs/``. If the property has digitized maps or tables that a person has
APPROVED in the review queue, a "Before & After" section shows each beside its scan (see
``samples``). It never reads PDFs itself: that is the ``digitize`` module's job.

ARIS is Access Only, so the pack is marked internal; get the Province's permission
before sending anything derived from ARIS to a prospect.
"""

from __future__ import annotations

from pathlib import Path

import geopandas as gpd
import pandas as pd
from docx import Document

from gossan.core.licensing import derived_license
from gossan.core.project import Project
from gossan.harvest.bc_provincial.aris_pdf import RAW_DIR as ARIS_PDF_DIR
from gossan.leads.generator import DERIVED_LICENSE, LEADS_LAYER, is_us_leads
from gossan.leads.samples import add_before_after
from gossan.leads.scoring import ARIS_LAYER, MissingLayerError, read_layer
from gossan.report.licensing import project_includes_aris


class LeadNotFoundError(Exception):
    """No lead with that id."""


def reports_for_lead(lead: gpd.GeoSeries, aris: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    """ARIS reports whose location falls inside the lead's tenure, oldest first."""
    inside = aris[aris.intersects(lead.geometry)]
    return inside.sort_values(["REPORT_YEAR", "ARIS_NUMBER"]).reset_index(drop=True)


def _drilling_text(lead: gpd.GeoSeries) -> str:
    """One line on the drilling flag; a permit signal, not proof that a rig is on site."""
    flag = lead.get("drilling_flag", "not assessed")
    if flag in ("not assessed", "none"):
        return "none found" if flag == "none" else "not assessed (harvest drilling_notices)"
    notes = [f"{flag}", f"notices {lead['drilling_notice_numbers']}"]
    if lead["drilling_operator"]:
        notes.append(str(lead["drilling_operator"]))
    if lead["drilling_latest_filed"]:
        notes.append(f"latest filed {lead['drilling_latest_filed']}")
    notes.append(f"+{lead['drilling_points']:g} points; permit signal only")
    return "; ".join(notes)


def build_pack(project: Project, lead_id: str) -> Path:
    """Write ``outputs/lead_<id>_pack.docx`` for a lead from the ``leads`` layer."""
    try:
        leads = read_layer(project, LEADS_LAYER)
    except MissingLayerError as e:
        raise MissingLayerError("no leads layer yet; run 'gossan leads run' first") from e
    match = leads[leads["lead_id"] == str(lead_id)]
    if match.empty:
        raise LeadNotFoundError(f"lead {lead_id!r} not found; see outputs/leads_ranked.csv")
    lead = match.iloc[0]
    if is_us_leads(leads):
        return _build_us_pack(project, lead, leads.crs)
    aris = read_layer(project, ARIS_LAYER).to_crs(leads.crs)
    reports = reports_for_lead(lead, aris)

    doc = Document()
    doc.add_heading(f"Lead {lead['lead_id']}: {lead['claim_name']}", level=0)
    doc.add_paragraph(
        "INTERNAL DRAFT. Contains ARIS data (Access Only): do not send to a prospect "
        "without the Province's written permission."
    )

    doc.add_heading("Property", level=1)
    fields = [
        ("Tenure type", lead["tenure_type"]),
        ("Owner", lead["owner"]),
        ("Issued", lead["issue_date"]),
        ("Good to", lead["good_to_date"]),
        ("Area (ha)", lead["area_ha"]),
        ("Lead rank / score", f"{lead['rank']} / {lead['score']}"),
        ("Recency points", lead["recency_points"]),
        ("Report-count points", lead["count_points"]),
        ("Legacy-report points", lead["legacy_points"]),
        ("Drilling notice (BC Notice of Work)", _drilling_text(lead)),
    ]
    table = doc.add_table(rows=0, cols=2, style="Light List")
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, "" if pd.isna(value) else str(value)

    doc.add_heading("Assessment reports on the property", level=1)
    if reports.empty:
        doc.add_paragraph("No ARIS reports are located on this tenure.")
    else:
        doc.add_paragraph(
            f"{len(reports)} report(s), {int(reports['REPORT_YEAR'].min())} to "
            f"{int(reports['REPORT_YEAR'].max())}. A report's location is a single point; "
            "it may cover ground beyond this tenure."
        )
        table = doc.add_table(rows=1, cols=6, style="Light List")
        for cell, head in zip(
            table.rows[0].cells,
            ["ARIS", "Year", "Title", "Operator", "Work", "PDF in project"],
            strict=True,
        ):
            cell.text = head
        for r in reports.itertuples():
            number = int(r.ARIS_NUMBER)
            have_pdf = (project.raw_dir / ARIS_PDF_DIR / f"ARIS_{number:05d}.pdf").exists()
            row = table.add_row().cells
            values = [
                str(number),
                str(int(r.REPORT_YEAR)),
                r.ASSESSMENT_REPORT_TITLE,
                r.OPERATOR_NAMES,
                r.GENERAL_WORK_TYPES,
                "yes" if have_pdf else "no",
            ]
            for cell, value in zip(row, values, strict=True):
                cell.text = "" if pd.isna(value) else str(value)

    path = project.outputs_dir / f"lead_{lead['lead_id']}_pack.docx"
    report_numbers = {str(int(n)) for n in reports["ARIS_NUMBER"]}
    samples = add_before_after(
        doc, project, lead.geometry, report_numbers, project.outputs_dir / f"{path.stem}_assets"
    )

    doc.add_heading("Next steps", level=1)
    doc.add_paragraph(
        "Approved samples above come from reports a person has checked. To add more, digitize "
        "further maps or tables from the reports listed and approve them with 'gossan review'."
        if samples.n_maps + samples.n_tables
        else "Pick two or three reports with maps or drill tables, run 'gossan digitize' on them "
        "and approve the results with 'gossan review' to fill the Before & After section."
    )

    project.outputs_dir.mkdir(exist_ok=True)
    doc.save(path)
    project.provenance.log(
        dataset="lead_pack",
        source="gossan",
        identifier=f"outputs/{path.name}",
        url=None,
        license=DERIVED_LICENSE,
        operation="build_lead_pack",
        parameters={
            "lead_id": str(lead["lead_id"]),
            "n_reports": len(reports),
            "n_approved_maps": samples.n_maps,
            "n_approved_tables": samples.n_tables,
            "n_pending_items": samples.n_pending,
        },
    )
    return path


def _build_us_pack(project: Project, lead: gpd.GeoSeries, crs: object) -> Path:
    """One-claim summary for a US lead: the claim, its spatial score, and nearby MRDS occurrences."""
    includes_aris = project_includes_aris(project)
    occ = read_layer(project, "minfile").to_crs(crs)
    near = occ[occ.distance(lead.geometry) <= 2000].copy()
    near["distance_m"] = near.distance(lead.geometry).round(0)
    near = near.sort_values("distance_m").head(25)

    doc = Document()
    doc.add_heading(f"Lead {lead['lead_id']}: {lead['claim_name']}", level=0)
    if includes_aris:
        doc.add_paragraph(
            "INTERNAL DRAFT. This project holds ARIS data (Access Only): do not send to a prospect "
            "without the Province's written permission."
        )
    else:
        doc.add_paragraph("DRAFT. Machine-generated: for geologist review before use.")
    doc.add_paragraph(
        "BLM claims are resolved to quarter-section geocodes, not surveyed claim boundaries, and "
        "MRDS is a historical inventory (not systematically updated since 2011). The score is spatial "
        "only: BLM MLRS has no issue or expiry dates, so recency is not scored."
    )
    doc.add_heading("Claim", level=1)
    fields = [
        ("Type", lead["tenure_type"]), ("Area (ha)", lead["area_ha"]),
        ("Lead rank / score", f"{lead['rank']} / {lead['score']}"),
        ("Target overlap (% of claim)", lead["target_overlap_pct"]),
        ("Target overlap points", lead["overlap_points"]),
        ("Nearest MRDS occurrence (m)", lead["nearest_occurrence_m"]),
        ("MRDS occurrences on the claim", lead["occurrences_inside"]),
        ("Proximity points", lead["proximity_points"]),
    ]  # fmt: skip
    table = doc.add_table(rows=0, cols=2, style="Light List")
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, "" if pd.isna(value) else str(value)

    doc.add_heading("MRDS occurrences within 2 km", level=1)
    if near.empty:
        doc.add_paragraph("No MRDS occurrences within 2 km of this claim.")
    else:
        table = doc.add_table(rows=1, cols=3, style="Light List")
        for cell, head in zip(table.rows[0].cells, ["Site", "Distance (m)", "Status"], strict=True):
            cell.text = head
        for r in near.itertuples():
            row = table.add_row().cells
            values = [getattr(r, "MINFILE_NAME1", ""), r.distance_m, getattr(r, "STATUS_DESCRIPTION", "")]
            for cell, value in zip(row, values, strict=True):
                cell.text = "" if pd.isna(value) else str(value)

    project.outputs_dir.mkdir(exist_ok=True)
    path = project.outputs_dir / f"lead_{lead['lead_id']}_pack.docx"
    doc.save(path)
    project.provenance.log(
        dataset="lead_pack",
        source="gossan",
        identifier=f"outputs/{path.name}",
        url=None,
        license=derived_license(project.provenance.read(), includes_aris=includes_aris),
        operation="build_lead_pack",
        parameters={"lead_id": str(lead["lead_id"]), "scoring": "us_spatial", "n_occurrences_listed": len(near)},
    )
    return path
