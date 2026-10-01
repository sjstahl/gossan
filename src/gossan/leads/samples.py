"""The "Before & After" section of a lead's sample pack.

For the ARIS reports on a property, this takes the review-queue items a person has
APPROVED (maps and tables) and shows each one next to the scan it came from:

* map: the rendered scan page beside the georeferenced result placed on the ground, with
  the lead's tenure outlined, so a prospect can see the old paper map lined up with today's
  claims;
* table: the scanned page beside a snippet of the approved, parsed data.

Pending and rejected items are never shown. Images are saved as files in
``outputs/lead_<id>_pack_assets/`` next to the Word document so the pack can be rebuilt and
checked. The images come from ARIS reports (Access Only): the pack stays internal.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import geopandas as gpd
import matplotlib
import rasterio
from docx.shared import Inches
from PIL import Image

from gossan.core.project import Project
from gossan.digitize.review import ReviewItem, approved_items, read_queue

matplotlib.use("Agg")
import matplotlib.pyplot as plt

THUMB_PX = 1000
SNIPPET_ROWS = 8
SNIPPET_COLS = 8
PICTURE_WIDTH = Inches(3.1)
OUTLINE = "#c0007a"


@dataclass
class SampleSummary:
    n_maps: int = 0
    n_tables: int = 0
    n_pending: int = 0
    missing: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        self.missing = self.missing or []


def thumbnail(src: Path, dest: Path) -> Path | None:
    """A PNG of ``src`` no larger than ``THUMB_PX`` on a side; None if it is not there."""
    if not src.is_file():
        return None
    dest.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as img:
        img = img.convert("RGB")
        img.thumbnail((THUMB_PX, THUMB_PX))
        img.save(dest)
    return dest


def georef_thumbnail(tif: Path, outline: gpd.GeoSeries | None, dest: Path) -> Path | None:
    """The georeferenced map drawn on its coordinates with ``outline`` (a tenure) over it."""
    if not tif.is_file():
        return None
    with rasterio.open(tif) as ds:
        scale = min(1.0, 1400 / max(ds.width, ds.height))
        shape = (ds.count, max(1, int(ds.height * scale)), max(1, int(ds.width * scale)))
        data = ds.read(out_shape=shape).transpose(1, 2, 0)
        bounds = ds.bounds
        crs = ds.crs
    if data.shape[2] == 2:  # gray + alpha
        data = data[:, :, [0, 0, 0, 1]]
    fig, ax = plt.subplots(figsize=(6, 4.2))
    ax.imshow(data, extent=(bounds.left, bounds.right, bounds.bottom, bounds.top))
    if outline is not None and len(outline):
        outline.to_crs(crs).boundary.plot(ax=ax, color=OUTLINE, linewidth=1.8)
    ax.set_xlim(bounds.left, bounds.right)
    ax.set_ylim(bounds.bottom, bounds.top)
    ax.ticklabel_format(style="plain", useOffset=False)
    ax.tick_params(labelsize=6, labelrotation=20)
    ax.set_xlabel(str(crs), fontsize=6)
    fig.tight_layout()
    dest.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(dest, dpi=110)
    plt.close(fig)
    return dest


def _latest_params(project: Project, operations: set[str], *, output: str = "", identifier: str = "") -> dict:
    """Parameters of the latest provenance record of these operations for an output file or id."""
    found: dict = {}
    for rec in project.provenance.read():
        if rec.operation in operations and (
            (output and rec.parameters.get("output") == output)
            or (identifier and rec.identifier == identifier)
        ):
            found = rec.parameters
    return found


def _sort_key(item: ReviewItem) -> tuple[int, int]:
    return int(item.report), int(item.page) if item.page.isdigit() else 0


def add_before_after(
    doc, project: Project, lead_geometry, report_numbers: set[str], assets: Path
) -> SampleSummary:
    """Append the section to ``doc``. ``lead_geometry`` is the tenure polygon (project CRS)."""
    summary = SampleSummary()
    items = [i for i in approved_items(project, report_numbers) if i.type in {"map", "table"}]
    summary.n_pending = sum(
        1 for i in read_queue(project) if i.status == "PENDING" and i.report in report_numbers
    )
    doc.add_heading("Before & After: digitized samples", level=1)
    if not items:
        note = "No approved digitized samples yet for this property."
        if summary.n_pending:
            note += f" {summary.n_pending} item(s) are waiting for review."
        doc.add_paragraph(note)
        return summary

    outline = gpd.GeoSeries([lead_geometry], crs=project.config.crs)
    for item in sorted(items, key=_sort_key):
        where = f"ARIS {item.report}" + (f", page {item.page}" if item.page else "")
        doc.add_heading(f"{where}: {'georeferenced map' if item.type == 'map' else 'extracted table'}", level=2)
        stem = f"{int(item.report):05d}_p{item.page or 0}_{item.type}"
        before = thumbnail(project.root / item.image, assets / f"{stem}_before.png") if item.image else None
        grid = doc.add_table(rows=2, cols=2)
        grid.rows[0].cells[0].text, grid.rows[0].cells[1].text = "Before: the scan", ""
        left, right = grid.rows[1].cells
        _put_picture(left, before, summary, f"{where} scan image")

        if item.type == "map":
            grid.rows[0].cells[1].text = "After: georeferenced, tenure outlined"
            after = georef_thumbnail(project.root / item.source, outline, assets / f"{stem}_after.png")
            _put_picture(right, after, summary, f"{where} georeferenced map")
            info = _latest_params(project, {"georef_auto", "georef_manual"}, output=item.source)
            detail = (
                f"{info['n_gcps']} control points, RMS {info['rms_m']} m, {info.get('pixel_size_m', '?')} m pixels"
                if info
                else "georeferencing details not found in provenance"
            )
            summary.n_maps += 1
        else:
            grid.rows[0].cells[1].text = "After: parsed data (excerpt)"
            total_rows, total_cols = _put_snippet(right, project.root / item.source, summary, where)
            loaded = _latest_params(project, {"load_drillhole_table"}, identifier=item.item_id)
            detail = f"{total_rows} row(s), {total_cols} column(s) in full" + (
                f"; loaded into the {loaded['table']} table as {loaded['records']} record(s)" if loaded else ""
            )
            summary.n_tables += 1
        doc.add_paragraph(
            f"{detail}. Approved by {item.reviewed_by or 'unknown'} on {item.reviewed_at or 'unknown date'}."
            + (f" Reviewer's note: {item.decision_note}" if item.decision_note else "")
        )
    return summary


def _put_picture(cell, image: Path | None, summary: SampleSummary, label: str) -> None:
    if image is None:
        cell.text = f"({label} not found)"
        summary.missing.append(label)
        return
    cell.paragraphs[0].add_run().add_picture(str(image), width=PICTURE_WIDTH)


def _put_snippet(cell, csv_path: Path, summary: SampleSummary, label: str) -> tuple[int, int]:
    if not csv_path.is_file():
        cell.text = f"({label} data file not found)"
        summary.missing.append(f"{label} data")
        return 0, 0
    with csv_path.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    header, body = rows[0], rows[1:]
    shown_cols = min(len(header), SNIPPET_COLS)
    table = cell.add_table(rows=1, cols=shown_cols)
    table.style = "Light List"
    for c, head in zip(table.rows[0].cells, header[:shown_cols], strict=True):
        c.text = head
    for row in body[:SNIPPET_ROWS]:
        for c, value in zip(table.add_row().cells, (row + [""] * shown_cols)[:shown_cols], strict=True):
            c.text = value
    return len(body), len(header)
