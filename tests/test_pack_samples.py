import getpass
from datetime import date
from pathlib import Path

import numpy as np
import pytest
import rasterio
from docx import Document
from PIL import Image
from rasterio.transform import from_bounds
from typer.testing import CliRunner

from gossan.cli import app
from gossan.core.project import Project
from gossan.digitize import review
from gossan.leads.generator import run_leads
from gossan.leads.scoring import read_layer

runner = CliRunner()
AS_OF = date(2026, 10, 1)
ASSAY_CSV = "hole_id,from_m,to_m,Au (ppb),Cu (ppm)\nR-1,0,1.5,12,40\nR-1,1.5,3.0,<5,55\nR-2,0,2,8,61\n"


@pytest.fixture(autouse=True)
def reviewer(monkeypatch):
    monkeypatch.setattr(getpass, "getuser", lambda: "tester")


@pytest.fixture
def project(aoi_file, seed_leads_layers) -> Project:
    assert runner.invoke(app, ["new", "demo", "--aoi", str(aoi_file)]).exit_code == 0
    proj = Project(Path("demo"))
    seed_leads_layers(proj)
    run_leads(proj, AS_OF)  # lead 1002 has ARIS reports 501-504
    return proj


def make_image(project: Project, rel: str) -> str:
    path = project.root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (600, 400), (240, 235, 220)).save(path)
    return rel


def queue_map(project: Project, report: int, page: int, *, georef=True) -> review.ReviewItem:
    """A georeferenced map the way digitize leaves it: image, COG over the tenure, provenance, item."""
    lead = read_layer(project, "leads").set_index("lead_id").loc["1002"]
    minx, miny, maxx, maxy = lead.geometry.bounds
    tif_rel = f"rasters/aris_{report:05d}_p{page}.tif"
    if georef:
        (project.root / "rasters").mkdir(exist_ok=True)
        data = np.full((4, 60, 90), 200, dtype=np.uint8)
        data[3] = 255
        with rasterio.open(
            project.root / tif_rel, "w", driver="GTiff", width=90, height=60, count=4, dtype="uint8",
            crs=project.config.crs, transform=from_bounds(minx - 30, miny - 30, maxx + 30, maxy + 30, 90, 60),
        ) as ds:  # fmt: skip
            ds.write(data)
        project.provenance.log(
            dataset="aris_maps", source="QGIS Georeferencer points (user)", identifier=f"ARIS {report} p{page}",
            url=None, license="x", operation="georef_manual",
            parameters={"output": tif_rel, "n_gcps": 9, "rms_m": 0.4, "pixel_size_m": 0.5},
        )  # fmt: skip
    image = make_image(project, f"digitize/maps/{report:05d}_p{page}.png")
    return review.add_item(project, report, "map", page=page, source=tif_rel, image=image, note="n")


def queue_table(project: Project, report: int, page: int) -> review.ReviewItem:
    rel = f"digitize/tables/scanned/{report:05d}_p{page}.csv"
    (project.root / rel).parent.mkdir(parents=True, exist_ok=True)
    (project.root / rel).write_text(ASSAY_CSV, encoding="utf-8")
    project.provenance.log(
        dataset="aris_tables_scanned", source="Anthropic API (Claude vision)", identifier=f"ARIS {report} p{page}",
        url=None, license="x", operation="extract_scanned_table", parameters={"output": rel, "table_type": "assay"},
    )  # fmt: skip
    image = make_image(project, f"digitize/pages/{report:05d}_p{page}.png")
    return review.add_item(project, report, "table", page=page, source=rel, image=image, note="n")


def pack_text(project: Project) -> tuple[Document, str]:
    result = runner.invoke(app, ["leads", "pack", str(project.root), "1002"])
    assert result.exit_code == 0, result.output
    doc = Document(project.outputs_dir / "lead_1002_pack.docx")
    return doc, "\n".join(p.text for p in doc.paragraphs)


def approve(project: Project, item_id: str, note: str = "") -> None:
    result = runner.invoke(app, ["review", "approve", str(project.root), item_id, "--note", note])
    assert result.exit_code == 0, result.output


def test_pack_shows_approved_map_and_table_before_and_after(project):
    m = queue_map(project, 502, 3)
    t = queue_table(project, 503, 5)
    approve(project, m.item_id, "placement matches the road")
    approve(project, t.item_id)

    doc, text = pack_text(project)
    assert "Before & After: digitized samples" in text
    assert "ARIS 502, page 3: georeferenced map" in text
    assert "ARIS 503, page 5: extracted table" in text
    assert "9 control points, RMS 0.4 m, 0.5 m pixels" in text
    assert "Approved by tester on 20" in text and "placement matches the road" in text
    assert "loaded into the assay table as 6 record(s)" in text and "3 row(s), 5 column(s) in full" in text

    assert len(doc.inline_shapes) == 3  # map: scan + georeferenced; table: scan
    snippet = next(t for outer in doc.tables for row in outer.rows for cell in row.cells for t in cell.tables)
    assert [c.text for c in snippet.rows[0].cells] == ["hole_id", "from_m", "to_m", "Au (ppb)", "Cu (ppm)"]
    assert [c.text for c in snippet.rows[2].cells] == ["R-1", "1.5", "3.0", "<5", "55"]  # as approved, flag kept

    assets = project.outputs_dir / "lead_1002_pack_assets"
    assert sorted(p.name for p in assets.iterdir()) == [
        "00502_p3_map_after.png", "00502_p3_map_before.png", "00503_p5_table_before.png",
    ]  # fmt: skip
    rec = next(r for r in project.provenance.read() if r.operation == "build_lead_pack")
    assert (rec.parameters["n_approved_maps"], rec.parameters["n_approved_tables"]) == (1, 1)


def test_only_approved_items_for_this_property_appear(project):
    approved = queue_map(project, 502, 3)
    pending = queue_map(project, 503, 4)  # waiting for review
    rejected = queue_map(project, 504, 6)
    elsewhere = queue_map(project, 999, 2)  # a report that is not on this tenure
    approve(project, approved.item_id)
    runner.invoke(app, ["review", "reject", str(project.root), rejected.item_id])
    approve(project, elsewhere.item_id)

    doc, text = pack_text(project)
    assert "ARIS 502, page 3" in text
    assert "ARIS 503" not in text.split("Before & After")[1] and "ARIS 504" not in text.split("Before & After")[1]
    assert "ARIS 999" not in text
    assert len(doc.inline_shapes) == 2
    rec = next(r for r in project.provenance.read() if r.operation == "build_lead_pack")
    assert rec.parameters["n_pending_items"] == 1 and pending.status == "PENDING"


def test_without_approved_items_the_pack_says_so_and_counts_pending(project):
    queue_map(project, 502, 3)
    doc, text = pack_text(project)
    assert "No approved digitized samples yet for this property. 1 item(s) are waiting for review." in text
    assert len(doc.inline_shapes) == 0 and not (project.outputs_dir / "lead_1002_pack_assets").exists()
    assert "run 'gossan digitize' on them" in text


def test_missing_files_are_noted_not_fatal(project):
    m = queue_map(project, 502, 3, georef=False)  # the georeferenced raster was never written
    approve(project, m.item_id)
    (project.root / m.image).unlink()
    doc, text = pack_text(project)
    cells = [c.text for outer in doc.tables for row in outer.rows for c in row.cells]
    assert any("scan image not found" in c for c in cells)
    assert any("georeferenced map not found" in c for c in cells)
    assert "georeferencing details not found in provenance" in text
    assert len(doc.inline_shapes) == 0


def test_pack_is_rebuilt_with_new_approvals(project):
    first = queue_map(project, 502, 3)
    approve(project, first.item_id)
    _, text = pack_text(project)
    assert "ARIS 503" not in text.split("Before & After")[1]
    second = queue_table(project, 503, 5)
    approve(project, second.item_id)
    doc, text = pack_text(project)
    assert "ARIS 503, page 5: extracted table" in text and len(doc.inline_shapes) == 3
