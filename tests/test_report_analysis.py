"""gossan report --kind analysis: one set of blocks, rendered to PDF and Word."""

from __future__ import annotations

import os
from pathlib import Path

import geopandas as gpd
import pymupdf
import pytest
from analysis_project import (
    ACCESS_ONLY_BC,
    CRS,
    EAST,
    X0,
    Y0,
    H,
    W,
    _write,
    build_analysis_project,
    runner,
)
from docx import Document
from shapely.geometry import LineString, box

from gossan.cli import app
from gossan.core.project import Project
from gossan.report.analysis_content import (
    Bullets,
    Figure,
    Heading,
    Table,
    build_content,
    load,
    pval,
    skill_word,
)
from gossan.report.analysis_render import CODE, _rl_text, banner_text


def full_project(tmp_path) -> Project:
    """The synthetic project plus tenure, a park and Access Only roads, with the target model and
    every analysis step already run."""
    project = build_analysis_project(tmp_path)
    _write(project, "mto_tenure", gpd.GeoDataFrame(
        {"TENURE_NUMBER_ID": [1], "TENURE_TYPE_DESCRIPTION": ["Mineral"], "OWNER_NAME": ["ACME LTD"], "GOOD_TO_DATE": ["2099-01-31"]},
        geometry=[box(EAST, Y0, X0 + W, Y0 + 1000)], crs=CRS))  # fmt: skip
    _write(project, "parks_protected", gpd.GeoDataFrame(
        {"PROTECTED_LANDS_NAME": ["TEST PARK"]}, geometry=[box(X0, Y0, X0 + 1000, Y0 + 1000)], crs=CRS))  # fmt: skip
    _write(project, "roads", gpd.GeoDataFrame(
        {"ROAD_CLASS": ["resource"]}, geometry=[LineString([(X0, Y0 + 2100), (X0 + W, Y0 + 2100)])], crs=CRS), ACCESS_ONLY_BC)  # fmt: skip
    assert runner.invoke(app, ["target", "demo", "--n-bags", "4", "--folds", "3", "--min-target-ha", "0"]).exit_code == 0
    result = runner.invoke(app, ["analyze", "demo"])
    assert result.exit_code == 0, result.output
    return Project(project.root.resolve())  # absolute, so tests can run from any directory


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """Build the project and both reports once for the module."""
    tmp = tmp_path_factory.mktemp("report")
    cwd = Path.cwd()
    os.chdir(tmp)
    try:
        project = full_project(tmp)
        result = runner.invoke(app, ["report", "demo", "--kind", "analysis"])
        assert result.exit_code == 0, result.output
        yield project, result
    finally:
        os.chdir(cwd)


# --- pure helpers -----------------------------------------------------------------------------


def test_skill_wording_follows_the_bands():
    assert skill_word(0.5) == "no clear skill"
    assert skill_word(0.6) == "modest skill"
    assert skill_word(0.7) == "moderate skill"
    assert skill_word(0.9) == "good skill"


def test_pvalue_and_code_spans():
    assert pval(0.0004) == "<0.001" and pval(0.0421) == "0.042"
    assert _rl_text("run `gossan analyze` <now> & go") == "run <font face='Courier'>gossan analyze</font> &lt;now&gt; &amp; go"
    assert [m.group(1) for m in CODE.finditer("a `b` c `d`")] == ["b", "d"]


# --- the built reports ----------------------------------------------------------------------


def test_both_files_are_written_and_logged(built):
    project, result = built
    out = project.root / "outputs"
    assert (out / "analysis_report.pdf").stat().st_size > 20_000 and (out / "analysis_report.docx").stat().st_size > 20_000
    assert "analysis_report.pdf" in result.output and "analysis_report.docx" in result.output
    records = {r.dataset: r for r in project.provenance.read() if r.operation == "build_analysis_report"}
    assert set(records) == {"report_analysis_pdf", "report_analysis_docx"}
    for r in records.values():
        assert "Access Only" in r.license and "roads" in r.license
        assert r.parameters["access_only_marking"] is True and "roads" in r.parameters["access_only_datasets"]
        assert r.parameters["sections"][:3] == ["summary", "data", "geology"] and len(r.parameters["sha256"]) == 64
    assert records["report_analysis_pdf"].parameters["pages"] >= 5


def test_pdf_content_and_marking(built):
    project, _ = built
    with pymupdf.open(project.root / "outputs" / "analysis_report.pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        footers = [page.get_text().count("INTERNAL DRAFT | Contains Access Only data") for page in pdf]
    for heading in ("Summary", "Read this first", "Data compiled", "Geology and known occurrences", "Stream-sediment geochemistry",
                    "Targets", "Tenure activity", "Target decision table", "Method and limits", "Data sources and licences"):  # fmt: skip
        assert heading in text, heading
    assert all(footers)  # every page is marked
    assert "INTERNAL DRAFT - NOT FOR DISTRIBUTION" in text and "roads" in text.split("written permission")[0]
    assert "Aldridge Formation" in text  # the planted enriched unit
    assert "Only 12 metallic occurrences are available" in text  # few-occurrence caveat
    assert "S000" in text  # the planted geochemical anomaly is a sample to follow up
    assert "equally explored ground" in text
    assert "TEST PARK" in text or "ground status" in text.lower()


def test_word_file_matches_the_pdf_structure(built):
    project, _ = built
    doc = Document(project.root / "outputs" / "analysis_report.docx")
    headings = [p.text for p in doc.paragraphs if p.style.name in {"Heading 1", "Title"}]
    assert headings[0] == "demo: analysis report"
    assert headings[1:] == ["Summary", "Data compiled", "Geology and known occurrences", "Stream-sediment geochemistry",
                            "Targets", "Tenure activity", "Target decision table", "Method and limits", "Data sources and licences"]  # fmt: skip
    assert len(doc.inline_shapes) >= 5 and len(doc.tables) >= 8
    assert any(p.style.name == "Caption" and p.text.startswith("Figure 1.") for p in doc.paragraphs)
    assert "INTERNAL DRAFT" in doc.sections[0].footer.paragraphs[0].text
    first_row = [c.text for c in next(t for t in doc.tables if t.rows[0].cells[0].text == "Dataset").rows[0].cells]
    assert first_row == ["Dataset", "Features", "Coverage", "Licence"]
    # `code` spans become monospace runs, not literal backticks.
    assert "`" not in "\n".join(p.text for p in doc.paragraphs)


def test_blocks_carry_the_numbers_and_figures(built, tmp_path):
    project, _ = built
    report = build_content(project, tmp_path / "figs")
    figures = [b for b in report.blocks if isinstance(b, Figure)]
    assert len(figures) >= 5 and all(f.path.exists() and f.path.stat().st_size > 5_000 for f in figures)
    summary = next(b for b in report.blocks if isinstance(b, Bullets))
    joined = " ".join(summary.items)
    assert "20 known MINFILE occurrences" in joined and "12 metallic" in joined
    assert "Aldridge Formation" in joined and "times the number expected" in joined
    decisions = next(b for b in report.blocks if isinstance(b, Table) and b.rows[0][:2] == ["#", "Target"])
    assert len(decisions.rows) > 3 and decisions.rows[0][4] == "Ground"
    assert [b.text for b in report.blocks if isinstance(b, Heading) and b.level == 1][:2] == ["Summary", "Data compiled"]


# --- degraded and open-data projects ----------------------------------------------------------


def test_report_on_a_project_with_no_analysis_says_what_to_run(tmp_path):
    aoi = tmp_path / "aoi.gpkg"
    gpd.GeoDataFrame({"id": [1]}, geometry=[box(X0, Y0, X0 + W, Y0 + H)], crs=CRS).to_file(aoi, driver="GPKG")
    assert runner.invoke(app, ["new", "bare", "--aoi", str(aoi)]).exit_code == 0
    result = runner.invoke(app, ["report", "bare", "--kind", "analysis", "--format", "pdf"])
    assert result.exit_code == 0, result.output
    assert not (Path("bare") / "outputs" / "analysis_report.docx").exists()
    with pymupdf.open(Path("bare") / "outputs" / "analysis_report.pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    for command in ("gossan analyze", "gossan target"):
        assert command in text
    assert "INTERNAL DRAFT" not in text and "Draft. Generated by Gossan" in text  # nothing Access Only: plain banner
    record = next(r for r in Project(Path("bare")).provenance.read() if r.operation == "build_analysis_report")
    assert record.parameters["access_only_marking"] is False


def test_open_data_projects_are_not_marked_internal(tmp_path):
    project = build_analysis_project(tmp_path, with_inventory=False)
    assert runner.invoke(app, ["analyze", "demo", "--steps", "inventory,geochem,occurrences"]).exit_code == 0
    assert runner.invoke(app, ["report", "demo", "--kind", "analysis", "--format", "pdf"]).exit_code == 0
    report = build_content(project, tmp_path / "figs")
    assert not report.marked and banner_text(report).startswith("Draft.")
    with pymupdf.open(Path("demo") / "outputs" / "analysis_report.pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    assert "INTERNAL DRAFT" not in text and "DRAFT | Machine-generated" in text
    record = next(r for r in project.provenance.read() if r.operation == "build_analysis_report")
    assert "Access Only" not in record.license and record.parameters["access_only_datasets"] == []


def test_stale_analysis_is_flagged(tmp_path):
    project = build_analysis_project(tmp_path)
    assert runner.invoke(app, ["analyze", "demo", "--steps", "inventory,geochem,occurrences"]).exit_code == 0
    assert load(project).stale == []
    old = project.data_gpkg.stat().st_mtime - 3600
    for name in ("inventory.csv", "geochem_summary.json"):
        os.utime(project.root / "analysis" / name, (old, old))
    assert load(project).stale == ["inventory.csv", "geochem_summary.json"]
    report = build_content(project, tmp_path / "figs")
    caveats = next(b for b in report.blocks if isinstance(b, Bullets) and any("older than the data" in i for i in b.items))
    assert any("gossan analyze" in i for i in caveats.items)


def test_unknown_kind_is_rejected(tmp_path):
    build_analysis_project(tmp_path)
    result = runner.invoke(app, ["report", "demo", "--kind", "essay"])
    assert result.exit_code != 0 and "kind" in result.output
