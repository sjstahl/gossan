from pathlib import Path

from docx import Document
from typer.testing import CliRunner

from gossan.cli import app
from gossan.core.project import Project

runner = CliRunner()


def test_phase1_done_when(aoi_file, wfs):
    """A claim polygon produces a project folder and a Word report with a MINFILE map and data sources."""
    assert runner.invoke(app, ["new", "demo", "--aoi", str(aoi_file)]).exit_code == 0
    assert runner.invoke(app, ["harvest", "demo"]).exit_code == 0
    result = runner.invoke(app, ["report", "demo"])
    assert result.exit_code == 0, result.output

    project = Project(Path("demo"))
    assert (project.outputs_dir / "map.png").stat().st_size > 0
    doc = Document(project.outputs_dir / "report.docx")
    text = "\n".join(p.text for p in doc.paragraphs)

    assert "demo" in text
    assert len(doc.inline_shapes) == 1  # the map
    assert "2 MINFILE occurrence(s)" in text
    sources = text.split("Data Sources", 1)[1]
    assert "minfile" in sources
    assert "pub:WHSE_MINERAL_TENURE.MINFIL_MINERAL_FILE" in sources
    assert "Open Government Licence - British Columbia" in sources

    operations = [r.operation for r in project.provenance.read()]
    assert operations[-2:] == ["render_map", "build_report"]


def test_report_without_harvest(aoi_file):
    runner.invoke(app, ["new", "demo", "--aoi", str(aoi_file)])
    result = runner.invoke(app, ["report", "demo"])
    assert result.exit_code == 0, result.output
    text = "\n".join(p.text for p in Document("demo/outputs/report.docx").paragraphs)
    assert "have not been harvested" in text
    assert "No external datasets" in text
