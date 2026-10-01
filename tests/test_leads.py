from datetime import date
from pathlib import Path

import geopandas as gpd
import pandas as pd
import pyogrio
import pytest
from docx import Document
from typer.testing import CliRunner

from gossan.cli import app
from gossan.core.project import Project
from gossan.leads.generator import run_leads
from gossan.leads.scoring import LeadsConfig, read_layer, recency_points, score_leads

runner = CliRunner()
AS_OF = date(2026, 10, 1)


@pytest.fixture
def project(aoi_file, seed_leads_layers) -> Project:
    assert runner.invoke(app, ["new", "demo", "--aoi", str(aoi_file)]).exit_code == 0
    proj = Project(Path("demo"))
    seed_leads_layers(proj)
    return proj


def scored(project: Project, config: LeadsConfig | None = None) -> gpd.GeoDataFrame:
    tenure = read_layer(project, "mto_tenure")
    aris = read_layer(project, "aris_index")
    return score_leads(tenure, aris, config or LeadsConfig(), AS_OF).set_index("lead_id")


def test_recency_tiers():
    tiers = [(1, 50), (3, 30), (5, 15)]
    assert recency_points(0.0, tiers) == 50
    assert recency_points(1.0, tiers) == 50  # boundary is inclusive
    assert recency_points(1.01, tiers) == 30
    assert recency_points(4.33, tiers) == 15
    assert recency_points(5.01, tiers) == 0
    assert recency_points(float("nan"), tiers) == 0  # no issue date


def test_scoring_math(project):
    leads = scored(project)
    # Hand-computed with default weights as of 2026-10-01:
    #   score = recency + 25 * min(n_reports, 10) / 10 + 25 * min(n_pre1990, 5) / 5
    expected = {
        "1003": (50, 7.5, 5, 62.5),  # issued 0.8 y ago; reports 1960, 1995, 2010
        "1001": (50, 0, 0, 50),  # issued 0.6 y ago; no reports
        "1004": (30, 2.5, 5, 37.5),  # issued 2.7 y ago; one 1980 report
        "1002": (0, 10, 15, 25),  # issued 2005; reports 1965, 1972, 1985, 1998
        "1007": (15, 0, 0, 15),  # issued 4.3 y ago
        "1008": (0, 0, 0, 0),  # no issue date
    }
    assert list(leads.index) == list(expected)  # ranked best first
    for lead_id, (recency, count, legacy, total) in expected.items():
        row = leads.loc[lead_id]
        assert (row.recency_points, row.count_points, row.legacy_points) == (recency, count, legacy)
        assert row.score == total
    assert list(leads["rank"]) == [1, 2, 3, 4, 5, 6]


def test_aris_stats_and_selection(project):
    leads = scored(project)
    row = leads.loc["1002"]
    assert (row.aris_count, row.aris_legacy_count) == (4, 3)
    assert (row.aris_earliest_year, row.aris_latest_year) == (1965, 1998)
    assert row.aris_numbers == "501; 502; 503; 504"
    # density = reports / tenure area (km2), from the projected polygon
    assert row.aris_per_km2 == pytest.approx(4 / (row.geometry.area / 1e6), abs=0.01)
    assert leads.loc["1001"].aris_count == 0 and leads.loc["1001"].aris_numbers == ""

    assert "1005" not in leads.index  # lapsed (good-to date before as-of)
    assert "1006" not in leads.index  # placer, not in tenure_types
    assert "541" not in "; ".join(leads.aris_numbers)  # ARIS point outside every tenure
    assert leads.loc["1007"].owner == "JOINT A; JOINT B"  # one row per tenure, owners merged
    assert leads.index.is_unique


def test_config_changes_selection_and_weights(project):
    config = LeadsConfig(active_only=False, tenure_types=["Mineral", "Placer"], legacy_points=0)
    leads = scored(project, config)
    assert {"1005", "1006"} <= set(leads.index)
    assert leads.loc["1005"].score == 5  # lapsed, issued 2010: 0 + 25*2/10 + 0
    assert leads.loc["1006"].score == 50
    assert leads.loc["1003"].score == 57.5  # legacy weight removed: 50 + 7.5


def test_config_read_from_project_yaml(project):
    project.config.parameters["leads"] = {"legacy_before_year": 1970}
    config = LeadsConfig.from_project(project)
    assert config.legacy_before_year == 1970
    with pytest.raises(ValueError):  # typos in parameters are errors, not silently ignored
        project.config.parameters["leads"] = {"legacy_befor_year": 1970}
        LeadsConfig.from_project(project)


def test_run_leads_writes_csv_and_layer(project):
    result = run_leads(project, AS_OF)
    assert len(result.leads) == 6

    csv = pd.read_csv(project.outputs_dir / "leads_ranked.csv", dtype={"lead_id": str})
    assert list(csv["lead_id"]) == ["1003", "1001", "1004", "1002", "1007", "1008"]
    assert list(csv["score"]) == [62.5, 50, 37.5, 25, 15, 0]
    assert list(csv["aris_count"]) == [3, 0, 1, 4, 0, 0]
    assert {"rank", "lead_id", "score", "aris_count", "aris_legacy_count"} <= set(csv.columns)
    assert csv.loc[csv["lead_id"] == "1002", "aris_numbers"].item() == "501; 502; 503; 504"

    assert "leads" in set(pyogrio.list_layers(project.data_gpkg)[:, 0])
    layer = gpd.read_file(project.data_gpkg, layer="leads")
    assert layer.crs.to_epsg() == 32610
    assert (layer.geom_type == "Polygon").all()
    assert sorted(layer["lead_id"]) == sorted(csv["lead_id"].astype(str))
    # Input layers are untouched.
    assert {"mto_tenure", "aris_index"} <= set(pyogrio.list_layers(project.data_gpkg)[:, 0])

    ops = {r.operation: r for r in project.provenance.read()}
    csv_rec, layer_rec = ops["write_leads_csv"], ops["write_leads_layer"]
    assert csv_rec.parameters["as_of"] == "2026-10-01"
    assert csv_rec.parameters["n_leads"] == 6
    assert "Access Only" in csv_rec.license  # restricted data stays flagged downstream
    assert layer_rec.identifier == "leads"


def test_run_leads_is_repeatable(project):
    first = run_leads(project, AS_OF).leads.drop(columns="geometry")
    second = run_leads(project, AS_OF).leads.drop(columns="geometry")
    pd.testing.assert_frame_equal(first, second)
    layer = gpd.read_file(project.data_gpkg, layer="leads")
    assert len(layer) == 6  # layer replaced, not appended


def test_cli_leads_run(project):
    result = runner.invoke(app, ["leads", "run", "demo", "--as-of", "2026-10-01"])
    assert result.exit_code == 0, result.output
    assert "6 lead(s) scored" in result.output
    assert "1003" in result.output and "RECENT RICH" in result.output
    assert (project.outputs_dir / "leads_ranked.csv").is_file()


def test_cli_leads_run_errors(aoi_file, project):
    assert runner.invoke(app, ["leads", "run", "demo", "--as-of", "yesterday"]).exit_code != 0

    runner.invoke(app, ["new", "bare", "--aoi", str(aoi_file)])
    result = runner.invoke(app, ["leads", "run", "bare"])
    assert result.exit_code == 1
    assert "mto_tenure" in result.output and "gossan harvest" in result.output


def test_no_matching_tenure_gives_empty_outputs(project):
    config = LeadsConfig(tenure_types=["Coal"])
    assert scored(project, config).empty
    project.config.parameters["leads"] = {"tenure_types": ["Coal"]}
    result = run_leads(project, AS_OF)
    assert result.leads.empty
    assert pd.read_csv(project.outputs_dir / "leads_ranked.csv").empty


def test_pack_lists_property_and_reports(project):
    run_leads(project, AS_OF)
    pdf_dir = project.raw_dir / "aris_pdfs"
    pdf_dir.mkdir(parents=True)
    (pdf_dir / "ARIS_00502.pdf").write_bytes(b"%PDF-1.4")  # one report already downloaded

    result = runner.invoke(app, ["leads", "pack", "demo", "1002"])
    assert result.exit_code == 0, result.output
    path = project.outputs_dir / "lead_1002_pack.docx"
    assert path.is_file()

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Lead 1002: OLD RICH" in text
    assert "INTERNAL DRAFT" in text and "Access Only" in text
    assert "4 report(s), 1965 to 1998" in text
    props = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    assert props["Owner"] == "OLD OWNER INC"
    assert props["Lead rank / score"] == "4 / 25.0"
    rows = [[c.text for c in row.cells] for row in doc.tables[1].rows[1:]]
    assert [r[0] for r in rows] == ["501", "502", "503", "504"]  # oldest first
    assert [r[1] for r in rows] == ["1965", "1972", "1985", "1998"]
    assert [r[5] for r in rows] == ["no", "yes", "no", "no"]  # PDF already in raw/aris_pdfs
    assert rows[0][2] == "Report 501"

    record = project.provenance.read()[-1]
    assert record.operation == "build_lead_pack"
    assert record.parameters == {"lead_id": "1002", "n_reports": 4}


def test_pack_for_lead_without_reports(project):
    run_leads(project, AS_OF)
    assert runner.invoke(app, ["leads", "pack", "demo", "1001"]).exit_code == 0
    text = "\n".join(p.text for p in Document(project.outputs_dir / "lead_1001_pack.docx").paragraphs)
    assert "No ARIS reports are located on this tenure." in text


def test_pack_errors(project):
    result = runner.invoke(app, ["leads", "pack", "demo", "1002"])  # before 'leads run'
    assert result.exit_code == 1 and "gossan leads run" in result.output
    run_leads(project, AS_OF)
    result = runner.invoke(app, ["leads", "pack", "demo", "9999"])
    assert result.exit_code == 1 and "not found" in result.output
    assert not list(project.outputs_dir.glob("lead_*_pack.docx"))
