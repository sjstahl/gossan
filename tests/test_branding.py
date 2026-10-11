"""Branded Word reports: settings, logo, cover page, a company template, and the disclaimers branding cannot remove."""

from __future__ import annotations

from pathlib import Path

import pytest
from analysis_project import build_analysis_project, runner
from docx import Document
from docx.shared import RGBColor
from PIL import Image

from gossan.cli import app
from gossan.core.config import BrandingConfig, ProjectConfig
from gossan.core.project import Project
from gossan.report.analysis_render import (
    MARKED_FOOTER,
    PLAIN_FOOTER,
    build_analysis_report,
)
from gossan.report.branding import (
    Branding,
    BrandingError,
    resolve,
    tint,
    write_starter_template,
)

ACCENT = "#c2410c"


@pytest.fixture
def project(tmp_path, monkeypatch) -> Project:
    monkeypatch.chdir(tmp_path)
    return build_analysis_project(tmp_path, with_inventory=False)  # no restricted data: an unmarked report


@pytest.fixture
def marked(tmp_path, monkeypatch) -> Project:
    monkeypatch.chdir(tmp_path)
    return build_analysis_project(tmp_path)  # holds an Access Only layer: a marked report


def logo(path: Path, size=(300, 90)) -> Path:
    Image.new("RGB", size, "#c2410c").save(path)
    return path


def brand(project: Project, **settings) -> Project:
    project.config.branding = BrandingConfig(**settings)
    project.config.save(project.root)
    return Project(project.root)


def docx_of(project: Project) -> Document:
    build_analysis_report(project, "docx")
    return Document(str(project.outputs_dir / "analysis_report.docx"))


def footer_text(doc: Document, first: bool = False) -> str:
    section = doc.sections[0]
    part = section.first_page_footer if first else section.footer
    return " ".join(p.text for p in part.paragraphs)


def header_images(doc: Document) -> list[str]:
    return [d.get("descr") for d in doc.sections[0].header._element.xpath(".//wp:docPr")]


# --- settings ------------------------------------------------------------------------------------------------------------------


def test_unbranded_reports_look_as_before_and_do_not_depend_on_a_table_style(project):
    doc = docx_of(project)
    assert PLAIN_FOOTER in footer_text(doc) and "demo" in footer_text(doc)
    assert doc.core_properties.author == "Gossan" and doc.styles["Normal"].font.name == "Calibri"
    assert doc.tables and all(t._tbl.tblPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders") is not None for t in doc.tables)
    assert not any(h.text for h in doc.sections[0].header.paragraphs)  # no company, no header


def test_company_logo_colour_and_fonts_are_applied(project, tmp_path):
    p = brand(project, company="Acme Geoscience", logo=str(logo(tmp_path / "logo.png")), accent=ACCENT, font="Arial", heading_font="Georgia", contact="geo@acme.example")
    doc = docx_of(p)
    header = " ".join(x.text for x in doc.sections[0].header.paragraphs)
    assert "Acme Geoscience" in header and header_images(doc) == ["Acme Geoscience logo"]  # the logo has alt text
    assert doc.styles["Normal"].font.name == "Arial" and doc.styles["Heading 1"].font.name == "Georgia"
    assert doc.styles["Heading 1"].font.color.rgb == RGBColor.from_string("C2410C") and doc.styles["Title"].font.color.rgb == RGBColor.from_string("C2410C")
    assert "Acme Geoscience" in footer_text(doc) and doc.core_properties.author == "Acme Geoscience"
    fills = {c._tc.tcPr.xpath("./w:shd/@w:fill")[0] for t in doc.tables for c in t.rows[0].cells if c._tc.tcPr is not None and c._tc.tcPr.xpath("./w:shd/@w:fill")}
    assert tint("C2410C") in fills  # table headers take the light tint of the accent
    record = [r for r in p.provenance.read() if r.dataset == "report_analysis_docx"][-1]
    assert record.parameters["branding"] == {"company": "Acme Geoscience", "template": None, "cover_page": False, "logo": "logo.png"}


def test_custom_footer_text_replaces_the_draft_line_on_an_unmarked_report(project):
    doc = docx_of(brand(project, footer_text="Acme Geoscience | Confidential to the client"))
    assert "Confidential to the client" in footer_text(doc) and PLAIN_FOOTER not in footer_text(doc)


# --- the disclaimers ---------------------------------------------------------------------------------------------------------


def test_branding_cannot_remove_the_access_only_marking(marked, tmp_path):
    p = brand(marked, company="Acme", footer_text="Looks clean", logo=str(logo(tmp_path / "l.png")), cover_page=True)
    doc = docx_of(p)
    for first in (False, True):  # every page, the cover included
        text = footer_text(doc, first)
        assert MARKED_FOOTER in text and "Looks clean" in text
    body = "\n".join(par.text for par in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "INTERNAL DRAFT - NOT FOR DISTRIBUTION" in body  # the banner stays in the document, on the cover
    red = [r.font.color.rgb for par in doc.sections[0].footer.paragraphs for r in par.runs if r.font.color and r.font.color.type is not None]
    assert RGBColor.from_string("7B1D12") in red


def test_a_template_footer_keeps_its_text_and_the_marking_is_added_after_it(marked, tmp_path):
    tpl = Document()
    tpl.sections[0].footer.paragraphs[0].text = "Acme house footer"
    tpl.sections[0].header.paragraphs[0].text = "ACME HEADER"
    tpl.add_paragraph("Template body text that must not appear")
    tpl.save(tmp_path / "t.docx")
    doc = docx_of(brand(marked, template=str(tmp_path / "t.docx"), company="Acme"))
    parts = [p.text for p in doc.sections[0].footer.paragraphs if p.text]
    assert parts[0] == "Acme house footer" and MARKED_FOOTER in parts[1]
    assert [p.text for p in doc.sections[0].header.paragraphs if p.text] == ["ACME HEADER"]  # not duplicated with the company line
    assert "Template body text" not in "\n".join(p.text for p in doc.paragraphs)


# --- cover page -----------------------------------------------------------------------------------------------------------------


def test_cover_page_holds_logo_company_title_client_and_contact(project, tmp_path):
    project.config.client = "Northern Ridge Resources"
    project.config.save(project.root)
    p = brand(Project(project.root), company="Acme Geoscience", logo=str(logo(tmp_path / "logo.png")), cover_page=True, contact="geo@acme.example")
    doc = docx_of(p)
    texts = [x.text for x in doc.paragraphs[:8]]
    assert any(t == "Acme Geoscience" for t in texts) and any("Prepared for Northern Ridge Resources" in t for t in texts) and any("geo@acme.example" in t for t in texts)
    assert doc.paragraphs[0]._p.xpath(".//w:drawing") and doc.sections[0].different_first_page_header_footer
    assert any(p_._p.xpath('.//w:br[@w:type="page"]') for p_ in doc.paragraphs[:12])  # the cover ends in a page break
    assert not any(x.text for x in doc.sections[0].first_page_header.paragraphs)  # no running header on the cover


# --- templates ---------------------------------------------------------------------------------------------------------------


def test_a_template_missing_styles_still_builds_and_gets_working_substitutes(project, tmp_path):
    tpl = Document()
    for name in ("List Bullet", "Caption", "Heading 2"):
        tpl.styles.element.remove(tpl.styles[name].element)
    tpl.save(tmp_path / "bare.docx")
    doc = docx_of(brand(project, template=str(tmp_path / "bare.docx")))
    assert {"List Bullet", "Caption", "Heading 2"} <= {s.name for s in doc.styles}
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert bullets and all(p.text.startswith("•") for p in bullets)  # the style has no list in it, so each bullet carries its own sign


def test_the_starter_template_shows_the_styles_and_the_company_header(project, tmp_path):
    p = brand(project, company="Acme Geoscience", logo=str(logo(tmp_path / "logo.png")), accent=ACCENT, contact="geo@acme.example")
    path = write_starter_template(tmp_path / "t" / "template.docx", resolve(p))
    doc = Document(str(path))
    assert {"Title", "Heading 1", "Heading 2", "Normal", "Caption", "List Bullet"} <= {s.name for s in doc.styles}
    assert "Acme Geoscience" in " ".join(x.text for x in doc.sections[0].header.paragraphs) and header_images(doc) == ["Acme Geoscience logo"]
    assert "Acme Geoscience | geo@acme.example" in footer_text(doc)
    # and a report built on it keeps that header once, with the footer line after the company line
    built = docx_of(brand(p, template=str(path)))
    assert header_images(built) == ["Acme Geoscience logo"] and PLAIN_FOOTER in footer_text(built) and "Acme Geoscience | geo@acme.example" in footer_text(built)


def test_the_template_command_writes_edits_and_refuses_to_overwrite(project, tmp_path):
    brand(project, company="Acme")
    first = runner.invoke(app, ["template", "demo"])
    assert first.exit_code == 0 and (project.root / "branding" / "template.docx").is_file() and "branding.template" in first.output
    again = runner.invoke(app, ["template", "demo"])
    assert again.exit_code == 1 and "--overwrite" in again.output
    assert runner.invoke(app, ["template", "demo", "--overwrite", "--out", str(tmp_path / "x.docx")]).exit_code == 0 and (tmp_path / "x.docx").is_file()


# --- errors ------------------------------------------------------------------------------------------------------------------------


def test_a_missing_or_unreadable_logo_or_template_stops_the_report_with_a_clear_message(project, tmp_path):
    for field, value, message in (
        ("logo", "nope.png", "branding.logo 'nope.png' not found"),
        ("template", "nope.docx", "branding.template 'nope.docx' not found"),
    ):
        p = brand(project, **{field: value})
        with pytest.raises(BrandingError, match=message):
            resolve(p)
        result = runner.invoke(app, ["report", "demo", "--kind", "analysis", "--format", "docx"])
        assert result.exit_code == 1 and message in result.output
    (tmp_path / "logo.png").write_text("not an image")
    with pytest.raises(BrandingError, match="not a readable image"):
        resolve(brand(project, logo=str(tmp_path / "logo.png")))
    (tmp_path / "bad.docx").write_text("not a docx")
    with pytest.raises(BrandingError, match=r"cannot open branding.template bad.docx.*\.dotx"):
        from gossan.report.branding import new_document

        new_document(Branding(template=tmp_path / "bad.docx"))


def test_a_pdf_only_report_does_not_need_the_branding_files(project):
    brand(project, logo="nope.png")
    assert runner.invoke(app, ["report", "demo", "--kind", "analysis", "--format", "pdf"]).exit_code == 0  # Word branding only


def test_settings_are_validated_and_old_projects_load(project):
    with pytest.raises(ValueError, match="colour"):
        BrandingConfig(accent="orange")
    data = project.config.model_dump()
    data.pop("branding")
    assert ProjectConfig.model_validate(data).branding.accent == "#2a78d6"
    assert tint("000000", 0.5) == "808080" and tint("2A78D6", 1.0) == "FFFFFF"


def test_the_simple_site_report_is_branded_too(project, tmp_path):
    p = brand(project, company="Acme Geoscience", logo=str(logo(tmp_path / "logo.png")))
    assert runner.invoke(app, ["report", "demo", "--kind", "pitch", "--format", "docx"]).exit_code == 0
    doc = Document(str(p.outputs_dir / "report.docx"))
    assert "Acme Geoscience" in " ".join(x.text for x in doc.sections[0].header.paragraphs) and PLAIN_FOOTER in footer_text(doc)
